#!/usr/bin/env python3
"""Build xlsx comparison sheet from tours.json. All 64 tours, sorted by departure date."""
import json
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "tours.json"
OUT = "tour-deals-2026-0925-1011.xlsx"

WEEK = "一二三四五六日"

def main():
    with open(SRC, encoding="utf-8") as f:
        tours = json.load(f)

    print(f"tour count: {len(tours)}")

    # normalize missing note key
    for t in tours:
        t.setdefault("note", "")

    tours.sort(key=lambda t: (t["date"], t["agency"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "9月25日-10月11日 團體旅遊比價"

    headers = ["出發日期", "星期", "旅行社", "行程名稱", "目的地／天數",
               "團費（每人／起）", "團況", "備註", "行程連結"]
    ws.append(headers)

    for t in tours:
        d = date.fromisoformat(t["date"])
        ws.append([
            d.strftime("%Y-%m-%d"),
            WEEK[d.weekday()],
            t["agency"],
            t["name"],
            t["dest"],
            t["price"],
            t["status"],
            t["note"],
            t["url"],
        ])

    n = len(tours) + 1

    # style header
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    # hyperlinks on the link column
    link_font = Font(color="0563C1", underline="single")
    for row in range(2, n + 1):
        url = ws.cell(row=row, column=9).value
        if url:
            ws.cell(row=row, column=9).value = "查看行程"
            ws.cell(row=row, column=9).hyperlink = url
            ws.cell(row=row, column=9).font = link_font

    # zebra striping
    stripe = PatternFill("solid", fgColor="DDEBF7")
    for row in range(2, n + 1):
        if row % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).fill = stripe

    # highlight confirmed groups
    hot_fill = PatternFill("solid", fgColor="FFF2CC")
    ref_fill = PatternFill("solid", fgColor="D9D9D9")
    for row in range(2, n + 1):
        status = str(ws.cell(row=row, column=7).value)
        if "僅供價格參考" in status:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).fill = ref_fill
            ws.cell(row=row, column=7).font = Font(bold=True, color="595959")
        elif any(k in status for k in ("已成團", "保證出團")):
            ws.cell(row=row, column=7).fill = hot_fill
            ws.cell(row=row, column=7).font = Font(bold=True, color="7F6000")

    widths = [12, 6, 10, 52, 16, 15, 22, 26, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A1:I{n}"
    ws.freeze_panes = "A2"

    # notes sheet
    ws2 = wb.create_sheet("說明")
    notes = [
        "資料來源：鳳凰旅遊、喜鴻假期、五福旅遊、雄獅旅遊、品冠旅遊、山富旅遊、可樂旅遊、理想旅遊 官網",
        "查詢時間：2026-09-16（團況與機位為查詢當下即時狀態，會變動）",
        "團費皆為每人「起」價，多以雙人一室計價；單人房差各官網多未標示，需直接洽詢旅行社",
        "黃色標示 = 已成團／保證出團；灰色標示 = 使用者 9/14 取消的喜鴻「風華峴港5日」訂單，僅供價格參考（訂單總額 NT$44,541，未付訂金）",
        "9/25-9/28 為中秋連假，10/9-10/11 為雙十連假，連假團價較高且名額有限",
        "純研究，未報名、未付款、未填寫任何個人資料",
    ]
    for i, line in enumerate(notes, start=1):
        ws2.cell(row=i, column=1, value=line)
    ws2.column_dimensions["A"].width = 90

    wb.save(OUT)
    print(f"saved {OUT}")

if __name__ == "__main__":
    main()
